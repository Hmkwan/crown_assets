"""
统计报表路由模块
提供完善的数据统计和报表导出功能
"""
from flask import render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from app import db
from app.models import (
    User, Equipment, RepairOrder, SparePart, PartRequestOrder,
    EquipmentTransfer, EquipmentScrap, EquipmentLoan, 
    Department, EquipmentType, AssetCost, AssetLifecycle
)
from sqlalchemy import func, extract, case, and_, or_
from datetime import datetime, timedelta
from app.main import bp
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def check_report_permission():
    """检查报表权限"""
    if current_user.role != 'admin' and not getattr(current_user, 'can_view_reports', False):
        flash('您没有权限访问此页面', 'danger')
        return False
    return True


@bp.route('/statistics')
@login_required
def statistics_dashboard():
    """统计报表总览页面"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 基础统计
    stats = {
        'users': User.query.count(),
        'equipment': Equipment.query.count(),
        'departments': Department.query.count(),
        'equipment_types': EquipmentType.query.count(),
        'repair_orders': RepairOrder.query.count(),
        'part_requests': PartRequestOrder.query.count(),
        'transfers': EquipmentTransfer.query.count(),
        'loans': EquipmentLoan.query.count(),
        'scraps': EquipmentScrap.query.count(),
    }
    
    # 设备状态分布
    equipment_status = db.session.query(
        Equipment.status,
        func.count(Equipment.id).label('count')
    ).group_by(Equipment.status).all()
    
    # 按部门统计设备
    equipment_by_dept = db.session.query(
        Department.name.label('dept_name'),
        func.count(Equipment.id).label('count')
    ).join(Equipment, Equipment.department_id == Department.id
    ).group_by(Department.name).all()
    
    # 维修工单状态分布
    repair_status = db.session.query(
        RepairOrder.status,
        func.count(RepairOrder.id).label('count')
    ).group_by(RepairOrder.status).all()
    
    # 本月统计
    now = datetime.now()
    month_start = datetime(now.year, now.month, 1)
    
    monthly_stats = {
        'new_equipment': Equipment.query.filter(Equipment.purchase_date >= month_start).count(),
        'new_repair_orders': RepairOrder.query.filter(RepairOrder.created_date >= month_start).count(),
        'completed_repairs': RepairOrder.query.filter(
            RepairOrder.completed_date >= month_start,
            RepairOrder.status == 'completed'
        ).count(),
        'new_part_requests': PartRequestOrder.query.filter(
            PartRequestOrder.created_date >= month_start
        ).count(),
    }
    
    # 近12个月趋势数据
    months_data = []
    for i in range(11, -1, -1):
        target_date = now - timedelta(days=30*i)
        month_start = datetime(target_date.year, target_date.month, 1)
        if target_date.month == 12:
            month_end = datetime(target_date.year + 1, 1, 1)
        else:
            month_end = datetime(target_date.year, target_date.month + 1, 1)
        
        repair_count = RepairOrder.query.filter(
            and_(RepairOrder.created_date >= month_start, RepairOrder.created_date < month_end)
        ).count()
        
        months_data.append({
            'month': f"{target_date.year}-{target_date.month:02d}",
            'repair_orders': repair_count
        })
    
    return render_template('main/statistics_dashboard.html',
                         title='统计报表中心',
                         stats=stats,
                         equipment_status=equipment_status,
                         equipment_by_dept=equipment_by_dept,
                         repair_status=repair_status,
                         monthly_stats=monthly_stats,
                         months_data=months_data)


@bp.route('/statistics/equipment')
@login_required
def statistics_equipment():
    """设备统计报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 获取筛选参数
    dept_id = request.args.get('dept_id', type=int)
    type_id = request.args.get('type_id', type=int)
    status = request.args.get('status', '').strip()
    
    # 构建查询
    query = Equipment.query
    if dept_id:
        query = query.filter_by(department_id=dept_id)
    if type_id:
        query = query.filter_by(type_id=type_id)
    if status:
        query = query.filter_by(status=status)
    
    # 获取数据 - 按 id 降序排列（最新的设备在前）
    equipments = query.order_by(Equipment.id.desc()).all()
    
    # 统计汇总
    total_count = len(equipments)
    total_value = sum([float(eq.price or 0) for eq in equipments])
    
    # 按状态统计
    status_stats = db.session.query(
        Equipment.status,
        func.count(Equipment.id).label('count')
    ).filter(query.whereclause if hasattr(query, 'whereclause') else True
    ).group_by(Equipment.status).all()
    
    # 按类型统计
    type_stats = db.session.query(
        EquipmentType.name.label('type_name'),
        func.count(Equipment.id).label('count'),
        func.sum(Equipment.price).label('total_value')
    ).join(Equipment, Equipment.type_id == EquipmentType.id
    ).filter(query.whereclause if hasattr(query, 'whereclause') else True
    ).group_by(EquipmentType.name).all()
    
    # 获取部门和类型列表（用于筛选）
    departments = Department.query.order_by(Department.name).all()
    equipment_types = EquipmentType.query.order_by(EquipmentType.name).all()
    
    return render_template('main/statistics_equipment.html',
                         title='设备统计报表',
                         equipments=equipments,
                         total_count=total_count,
                         total_value=total_value,
                         status_stats=status_stats,
                         type_stats=type_stats,
                         departments=departments,
                         equipment_types=equipment_types,
                         selected_dept=dept_id,
                         selected_type=type_id,
                         selected_status=status)


@bp.route('/statistics/repair-orders')
@login_required
def statistics_repair_orders():
    """维修工单统计报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 获取筛选参数
    status_filter = request.args.get('status', '').strip()
    dept_id = request.args.get('dept_id', type=int)
    date_from = request.args.get('date_from', '').strip()
    date_to = request.args.get('date_to', '').strip()
    
    # 构建查询
    query = RepairOrder.query.join(Equipment).join(User, RepairOrder.requester_id == User.id)
    
    if status_filter:
        query = query.filter(RepairOrder.status == status_filter)
    if dept_id:
        query = query.filter(Equipment.department_id == dept_id)
    if date_from:
        query = query.filter(RepairOrder.created_date >= datetime.strptime(date_from, '%Y-%m-%d'))
    if date_to:
        query = query.filter(RepairOrder.created_date <= datetime.strptime(date_to + ' 23:59:59', '%Y-%m-%d %H:%M:%S'))
    
    # 获取数据
    repair_orders = query.order_by(RepairOrder.created_date.desc()).all()
    
    # 统计汇总
    total_count = len(repair_orders)
    completed_count = len([ro for ro in repair_orders if ro.status == 'completed'])
    pending_count = len([ro for ro in repair_orders if ro.status in ['submitted', 'in_progress']])
    
    # 计算平均处理时间（已完成的工单）
    completed_orders = [ro for ro in repair_orders if ro.status == 'completed' and ro.completed_date]
    if completed_orders:
        total_hours = sum([
            (ro.completed_date - ro.created_date).total_seconds() / 3600 
            for ro in completed_orders
        ])
        avg_hours = total_hours / len(completed_orders)
    else:
        avg_hours = 0
    
    # 按状态统计
    status_stats = db.session.query(
        RepairOrder.status,
        func.count(RepairOrder.id).label('count')
    ).group_by(RepairOrder.status).all()
    
    # 按技术员统计
    technician_stats = db.session.query(
        User.username,
        func.count(RepairOrder.id).label('total'),
        func.sum(case((RepairOrder.status == 'completed', 1), else_=0)).label('completed')
    ).join(RepairOrder, RepairOrder.technician_id == User.id
    ).group_by(User.id, User.username).all()
    
    # 获取部门列表
    departments = Department.query.order_by(Department.name).all()
    
    return render_template('main/statistics_repair_orders.html',
                         title='维修工单统计',
                         repair_orders=repair_orders,
                         total_count=total_count,
                         completed_count=completed_count,
                         pending_count=pending_count,
                         avg_hours=avg_hours,
                         status_stats=status_stats,
                         technician_stats=technician_stats,
                         departments=departments,
                         selected_status=status_filter,
                         selected_dept=dept_id,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/cost-analysis')
@login_required
def statistics_cost_analysis():
    """成本分析报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 设备总成本
    equipment_cost = db.session.query(
        func.sum(Equipment.price).label('total_purchase')
    ).scalar() or 0
    
    # 维修总成本
    repair_cost = db.session.query(
        func.sum(AssetCost.maintenance_cost).label('total_maintenance')
    ).scalar() or 0
    
    # 按部门统计成本
    dept_cost = db.session.query(
        Department.name.label('dept_name'),
        func.count(Equipment.id).label('equipment_count'),
        func.sum(Equipment.price).label('purchase_cost'),
        func.sum(AssetCost.maintenance_cost).label('maintenance_cost')
    ).join(Equipment, Equipment.department_id == Department.id
    ).outerjoin(AssetCost, AssetCost.equipment_id == Equipment.id
    ).group_by(Department.name).all()
    
    # 按类型统计成本
    type_cost = db.session.query(
        EquipmentType.name.label('type_name'),
        func.count(Equipment.id).label('equipment_count'),
        func.sum(Equipment.price).label('purchase_cost'),
        func.avg(Equipment.price).label('avg_price'),
        func.sum(AssetCost.maintenance_cost).label('maintenance_cost')
    ).join(Equipment, Equipment.type_id == EquipmentType.id
    ).outerjoin(AssetCost, AssetCost.equipment_id == Equipment.id
    ).group_by(EquipmentType.name).all()
    
    # 按年度统计
    year_cost = db.session.query(
        extract('year', Equipment.purchase_date).label('year'),
        func.count(Equipment.id).label('equipment_count'),
        func.sum(Equipment.price).label('purchase_cost')
    ).filter(Equipment.purchase_date.isnot(None)
    ).group_by(extract('year', Equipment.purchase_date)
    ).order_by(extract('year', Equipment.purchase_date).desc()).all()
    
    return render_template('main/statistics_cost_analysis.html',
                         title='成本分析报表',
                         equipment_cost=equipment_cost,
                         repair_cost=repair_cost,
                         dept_cost=dept_cost,
                         type_cost=type_cost,
                         year_cost=year_cost)


@bp.route('/statistics/spare-parts')
@login_required
def statistics_spare_parts():
    """配件统计报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 获取查询参数
    part_type = request.args.get('part_type', '')
    department_id = request.args.get('department_id', '')
    
    # 配件库存统计
    spare_parts_query = SparePart.query
    
    if department_id:
        spare_parts_query = spare_parts_query.filter_by(department_id=department_id)
    if part_type:
        spare_parts_query = spare_parts_query.filter(SparePart.name.like(f'%{part_type}%'))
    
    spare_parts = spare_parts_query.order_by(SparePart.id.desc()).all()
    
    # 库存汇总
    total_parts = len(spare_parts)
    total_value = sum([float(sp.price or 0) * (sp.stock_quantity or 0) for sp in spare_parts])
    low_stock_count = len([sp for sp in spare_parts if sp.stock_quantity is not None and sp.min_stock_level and sp.stock_quantity <= sp.min_stock_level])
    
    # 按部门统计
    dept_stats = db.session.query(
        Department.name,
        func.count(SparePart.id).label('count'),
        func.sum(SparePart.stock_quantity).label('total_quantity'),
        func.sum(SparePart.price * SparePart.stock_quantity).label('total_value')
    ).join(Department, Department.id == SparePart.department_id
    ).group_by(Department.name).all()
    
    # 配件申请统计
    request_status = request.args.get('request_status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    
    part_requests_query = PartRequestOrder.query
    
    if request_status:
        part_requests_query = part_requests_query.filter_by(status=request_status)
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, '%Y-%m-%d')
            part_requests_query = part_requests_query.filter(PartRequestOrder.created_date >= date_from_obj)
        except:
            pass
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, '%Y-%m-%d')
            part_requests_query = part_requests_query.filter(PartRequestOrder.created_date <= date_to_obj)
        except:
            pass
    
    part_requests = part_requests_query.order_by(PartRequestOrder.created_date.desc()).limit(50).all()
    
    # 申请单统计
    request_total = PartRequestOrder.query.count()
    request_pending = PartRequestOrder.query.filter_by(status='pending').count()
    request_approved = PartRequestOrder.query.filter_by(status='approved').count()
    request_completed = PartRequestOrder.query.filter_by(status='completed').count()
    
    # 按状态统计申请单
    request_status_stats = db.session.query(
        PartRequestOrder.status,
        func.count(PartRequestOrder.id).label('count')
    ).group_by(PartRequestOrder.status).all()
    
    # 近12个月申请趋势
    now = datetime.now()
    monthly_requests = []
    for i in range(11, -1, -1):
        target_date = now - timedelta(days=30*i)
        month_start = datetime(target_date.year, target_date.month, 1)
        if target_date.month == 12:
            month_end = datetime(target_date.year + 1, 1, 1)
        else:
            month_end = datetime(target_date.year, target_date.month + 1, 1)
        
        count = PartRequestOrder.query.filter(
            and_(PartRequestOrder.created_date >= month_start, PartRequestOrder.created_date < month_end)
        ).count()
        
        monthly_requests.append({
            'month': f"{target_date.year}-{target_date.month:02d}",
            'count': count
        })
    
    # 获取所有部门用于筛选
    departments = Department.query.order_by(Department.name).all()
    
    return render_template('main/statistics_spare_parts.html',
                         title='配件统计报表',
                         spare_parts=spare_parts,
                         total_parts=total_parts,
                         total_value=total_value,
                         low_stock_count=low_stock_count,
                         dept_stats=dept_stats,
                         part_requests=part_requests,
                         request_total=request_total,
                         request_pending=request_pending,
                         request_approved=request_approved,
                         request_completed=request_completed,
                         request_status_stats=request_status_stats,
                         monthly_requests=monthly_requests,
                         departments=departments,
                         # 查询参数
                         part_type=part_type,
                         department_id=department_id,
                         request_status=request_status,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/part-requests')
@login_required
def statistics_part_requests():
    """配件申请报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 获取查询参数
    status = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    
    # 查询配件申请
    query = PartRequestOrder.query
    
    if status:
        query = query.filter_by(status=status)
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, '%Y-%m-%d')
            query = query.filter(PartRequestOrder.created_date >= date_from_obj)
        except:
            pass
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, '%Y-%m-%d')
            query = query.filter(PartRequestOrder.created_date <= date_to_obj)
        except:
            pass
    
    part_requests = query.order_by(PartRequestOrder.created_date.desc()).all()
    
    # 统计汇总
    total_count = PartRequestOrder.query.count()
    pending_count = PartRequestOrder.query.filter_by(status='pending').count()
    approved_count = PartRequestOrder.query.filter_by(status='approved').count()
    completed_count = PartRequestOrder.query.filter_by(status='completed').count()
    rejected_count = PartRequestOrder.query.filter_by(status='rejected').count()
    
    # 按状态统计
    status_stats = db.session.query(
        PartRequestOrder.status,
        func.count(PartRequestOrder.id).label('count')
    ).group_by(PartRequestOrder.status).all()
    
    return render_template('main/statistics_part_requests.html',
                         title='配件申请报表',
                         part_requests=part_requests,
                         total_count=total_count,
                         pending_count=pending_count,
                         approved_count=approved_count,
                         completed_count=completed_count,
                         rejected_count=rejected_count,
                         status_stats=status_stats,
                         status=status,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/transfers')
@login_required
def statistics_transfers():
    """设备调拨报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 获取查询参数
    status = request.args.get('status', '')
    department_id = request.args.get('department_id', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    
    # 查询设备调拨
    query = EquipmentTransfer.query
    
    if status:
        query = query.filter_by(status=status)
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, '%Y-%m-%d')
            query = query.filter(EquipmentTransfer.created_date >= date_from_obj)
        except:
            pass
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, '%Y-%m-%d')
            query = query.filter(EquipmentTransfer.created_date <= date_to_obj)
        except:
            pass
    
    transfers = query.order_by(EquipmentTransfer.created_date.desc()).all()
    
    # 统计汇总
    total_count = EquipmentTransfer.query.count()
    pending_count = EquipmentTransfer.query.filter_by(status='pending').count()
    approved_count = EquipmentTransfer.query.filter_by(status='approved').count()
    completed_count = EquipmentTransfer.query.filter_by(status='completed').count()
    
    # 按状态统计
    status_stats = db.session.query(
        EquipmentTransfer.status,
        func.count(EquipmentTransfer.id).label('count')
    ).group_by(EquipmentTransfer.status).all()
    
    # 获取所有部门用于筛选
    departments = Department.query.order_by(Department.name).all()
    
    return render_template('main/statistics_transfers.html',
                         title='设备调拨报表',
                         transfers=transfers,
                         total_count=total_count,
                         pending_count=pending_count,
                         approved_count=approved_count,
                         completed_count=completed_count,
                         status_stats=status_stats,
                         departments=departments,
                         status=status,
                         department_id=department_id,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/asset-lifecycle')
@login_required
def statistics_asset_lifecycle():
    """资产生命周期报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 按事件类型统计
    event_stats = db.session.query(
        AssetLifecycle.event_type,
        func.count(AssetLifecycle.id).label('count'),
        func.sum(AssetLifecycle.cost_involved).label('total_cost')
    ).group_by(AssetLifecycle.event_type).all()
    
    # 最近事件
    recent_events = AssetLifecycle.query.order_by(
        AssetLifecycle.event_date.desc()
    ).limit(50).all()
    
    # 按月统计生命周期成本
    monthly_lifecycle = db.session.query(
        extract('year', AssetLifecycle.event_date).label('year'),
        extract('month', AssetLifecycle.event_date).label('month'),
        func.count(AssetLifecycle.id).label('event_count'),
        func.sum(AssetLifecycle.cost_involved).label('total_cost')
    ).filter(AssetLifecycle.event_date.isnot(None)
    ).group_by(
        extract('year', AssetLifecycle.event_date),
        extract('month', AssetLifecycle.event_date)
    ).order_by(
        extract('year', AssetLifecycle.event_date).desc(),
        extract('month', AssetLifecycle.event_date).desc()
    ).limit(12).all()
    
    return render_template('main/statistics_asset_lifecycle.html',
                         title='资产生命周期报表',
                         event_stats=event_stats,
                         recent_events=recent_events,
                         monthly_lifecycle=monthly_lifecycle)


@bp.route('/statistics/scraps')
@login_required
def statistics_scraps():
    """设备报废统计报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 筛选条件
    status = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    
    # 查询报废申请
    query = EquipmentScrap.query
    
    if status:
        query = query.filter_by(status=status)
    if date_from:
        query = query.filter(EquipmentScrap.created_date >= datetime.strptime(date_from, '%Y-%m-%d'))
    if date_to:
        query = query.filter(EquipmentScrap.created_date <= datetime.strptime(date_to + ' 23:59:59', '%Y-%m-%d %H:%M:%S'))
    
    scraps = query.order_by(EquipmentScrap.created_date.desc()).all()
    
    # 统计数据
    total_count = len(scraps)
    submitted_count = sum(1 for s in scraps if s.status in ['submitted', 'pending'])
    approved_count = sum(1 for s in scraps if s.status == 'approved')
    rejected_count = sum(1 for s in scraps if s.status in ['rejected', 'cancelled'])
    
    # 按状态分组统计
    status_stats = db.session.query(
        EquipmentScrap.status,
        func.count(EquipmentScrap.id).label('count')
    ).group_by(EquipmentScrap.status).all()
    
    return render_template('main/statistics_scraps.html',
                         title='设备报废报表',
                         scraps=scraps,
                         total_count=total_count,
                         submitted_count=submitted_count,
                         approved_count=approved_count,
                         rejected_count=rejected_count,
                         status_stats=status_stats,
                         status=status,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/loans')
@login_required
def statistics_loans():
    """设备借用统计报表"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 筛选条件
    status = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    
    # 查询借用记录
    query = EquipmentLoan.query
    
    if status:
        query = query.filter_by(status=status)
    if date_from:
        query = query.filter(EquipmentLoan.created_date >= datetime.strptime(date_from, '%Y-%m-%d'))
    if date_to:
        query = query.filter(EquipmentLoan.created_date <= datetime.strptime(date_to + ' 23:59:59', '%Y-%m-%d %H:%M:%S'))
    
    loans = query.order_by(EquipmentLoan.created_date.desc()).all()
    
    # 统计数据
    total_count = len(loans)
    submitted_count = sum(1 for l in loans if l.status == 'submitted')
    approved_count = sum(1 for l in loans if l.status == 'approved')
    borrowed_count = sum(1 for l in loans if l.status == 'borrowed')
    returned_count = sum(1 for l in loans if l.status == 'returned')
    rejected_count = sum(1 for l in loans if l.status in ['rejected', 'cancelled'])
    
    # 逾期统计
    overdue_loans = [l for l in loans if l.status == 'borrowed' and l.end_date and l.end_date < datetime.now()]
    overdue_count = len(overdue_loans)
    
    # 按状态分组统计
    status_stats = db.session.query(
        EquipmentLoan.status,
        func.count(EquipmentLoan.id).label('count')
    ).group_by(EquipmentLoan.status).all()
    
    return render_template('main/statistics_loans.html',
                         title='设备借用报表',
                         loans=loans,
                         total_count=total_count,
                         submitted_count=submitted_count,
                         approved_count=approved_count,
                         borrowed_count=borrowed_count,
                         returned_count=returned_count,
                         rejected_count=rejected_count,
                         overdue_count=overdue_count,
                         overdue_loans=overdue_loans,
                         status_stats=status_stats,
                         status=status,
                         date_from=date_from,
                         date_to=date_to)


@bp.route('/statistics/export/<report_type>')
@login_required
def export_statistics_xlsx(report_type):
    """导出统计报表为 XLSX 格式"""
    if not check_report_permission():
        return redirect(url_for('main.index'))
    
    # 创建工作簿
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认工作表
    
    # 定义样式
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    
    if report_type == 'equipment':
        # 设备报表
        ws = wb.create_sheet("设备清单")
        headers = ['设备ID', '设备名称', '设备类型', '品牌', '型号', '序列号', 
                  '所属部门', '位置', '状态', '采购价格', '购买日期']
        ws.append(headers)
        
        # 应用表头样式
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # 获取数据
        equipments = Equipment.query.order_by(Equipment.id).all()
        for eq in equipments:
            # department may be a string (legacy) or a relationship object; handle both
            dept_name = ''
            if eq.department:
                if isinstance(eq.department, str):
                    dept_name = eq.department
                else:
                    dept_name = getattr(eq.department, 'name', '')

            ws.append([
                eq.id,
                eq.name,
                eq.equipment_type.name if eq.equipment_type else '',
                eq.brand or '',
                eq.model or '',
                eq.serial_number,
                dept_name,
                eq.location or '',
                eq.status,
                float(eq.price or 0),
                eq.purchase_date.strftime('%Y-%m-%d') if eq.purchase_date else ''
            ])
        
        # 调整列宽
        column_widths = [10, 25, 15, 15, 15, 20, 15, 20, 12, 12, 12]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'设备报表_{timestamp}.xlsx'
    
    elif report_type == 'repair_orders':
        # 维修工单报表
        ws = wb.create_sheet("维修工单")
        headers = ['工单ID', '设备名称', '设备部门', '故障描述', '申请人', 
                  '技术员', '状态', '创建时间', '完成时间', '处理时长(小时)']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        repair_orders = RepairOrder.query.order_by(RepairOrder.created_date.desc()).all()
        for ro in repair_orders:
            duration = ''
            if ro.completed_date and ro.created_date:
                hours = (ro.completed_date - ro.created_date).total_seconds() / 3600
                duration = f"{hours:.1f}"
            
            ws.append([
                ro.id,
                ro.equipment.name if ro.equipment else '',
                # department may be a string (legacy) or a relationship object; handle both
                (ro.equipment.department if isinstance(ro.equipment.department, str) else getattr(ro.equipment.department, 'name', '')) if ro.equipment and ro.equipment.department else '',
                ro.description or '',
                ro.requester.real_name or ro.requester.username if ro.requester else '',
                ro.technician.real_name or ro.technician.username if ro.technician else '',
                ro.status,
                ro.created_date.strftime('%Y-%m-%d %H:%M:%S') if ro.created_date else '',
                ro.completed_date.strftime('%Y-%m-%d %H:%M:%S') if ro.completed_date else '',
                duration
            ])
        
        column_widths = [10, 25, 15, 35, 12, 12, 12, 20, 20, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'维修工单报表_{timestamp}.xlsx'
    
    elif report_type == 'spare_parts':
        # 配件库存报表
        ws = wb.create_sheet("配件库存")
        headers = ['配件ID', '配件名称', '配件编号', '库存数量', '单价', '总价', 
                  '最低库存', '所属部门', '位置', '购买日期', '是否公开']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        spare_parts = SparePart.query.order_by(SparePart.id).all()
        for sp in spare_parts:
            ws.append([
                sp.id,
                sp.name,
                sp.part_number or '',
                sp.stock_quantity or 0,
                float(sp.price or 0),
                float((sp.price or 0) * (sp.stock_quantity or 0)),
                sp.min_stock_level or '',
                sp.department or '',
                sp.location or '',
                sp.purchase_date.strftime('%Y-%m-%d') if sp.purchase_date else '',
                '是' if sp.is_public else '否'
            ])
        
        column_widths = [10, 25, 20, 12, 12, 12, 12, 15, 15, 12, 10]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'配件库存报表_{timestamp}.xlsx'
    
    elif report_type == 'cost_analysis':
        # 成本分析报表 - 多个工作表
        
        # 1. 按部门统计
        ws1 = wb.create_sheet("部门成本统计")
        headers1 = ['部门名称', '设备数量', '采购成本', '维修成本', '总成本']
        ws1.append(headers1)
        
        for col in range(1, len(headers1) + 1):
            cell = ws1.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        dept_cost = db.session.query(
            Department.name.label('dept_name'),
            func.count(Equipment.id).label('equipment_count'),
            func.sum(Equipment.price).label('purchase_cost'),
            func.sum(AssetCost.maintenance_cost).label('maintenance_cost')
        ).join(Equipment, Equipment.department_id == Department.id
        ).outerjoin(AssetCost, AssetCost.equipment_id == Equipment.id
        ).group_by(Department.name).all()
        
        for row in dept_cost:
            purchase = float(row.purchase_cost or 0)
            maintenance = float(row.maintenance_cost or 0)
            ws1.append([
                row.dept_name,
                row.equipment_count,
                purchase,
                maintenance,
                purchase + maintenance
            ])
        
        for i, width in enumerate([20, 12, 15, 15, 15], 1):
            ws1.column_dimensions[get_column_letter(i)].width = width
        
        # 2. 按类型统计
        ws2 = wb.create_sheet("类型成本统计")
        headers2 = ['设备类型', '数量', '采购成本', '平均价格', '维修成本']
        ws2.append(headers2)
        
        for col in range(1, len(headers2) + 1):
            cell = ws2.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        type_cost = db.session.query(
            EquipmentType.name.label('type_name'),
            func.count(Equipment.id).label('equipment_count'),
            func.sum(Equipment.price).label('purchase_cost'),
            func.avg(Equipment.price).label('avg_price'),
            func.sum(AssetCost.maintenance_cost).label('maintenance_cost')
        ).join(Equipment, Equipment.type_id == EquipmentType.id
        ).outerjoin(AssetCost, AssetCost.equipment_id == Equipment.id
        ).group_by(EquipmentType.name).all()
        
        for row in type_cost:
            ws2.append([
                row.type_name,
                row.equipment_count,
                float(row.purchase_cost or 0),
                float(row.avg_price or 0),
                float(row.maintenance_cost or 0)
            ])
        
        for i, width in enumerate([20, 12, 15, 15, 15], 1):
            ws2.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'成本分析报表_{timestamp}.xlsx'
    
    elif report_type == 'part_requests':
        # 配件申请报表
        ws = wb.create_sheet("配件申请")
        headers = ['申请ID', '配件名称', '配件编号', '数量', '单位', 
                  '申请人', '申请部门', '状态', '创建时间', '备注']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        part_requests = PartRequestOrder.query.order_by(PartRequestOrder.created_date.desc()).all()
        for pr in part_requests:
            ws.append([
                pr.id,
                pr.part_name or '',
                pr.part_number or '',
                pr.quantity or 0,
                pr.unit or '',
                pr.requester.real_name or pr.requester.username if pr.requester else '',
                pr.requester.department if pr.requester else '',
                pr.status,
                pr.created_date.strftime('%Y-%m-%d %H:%M:%S') if pr.created_date else '',
                pr.notes or ''
            ])
        
        column_widths = [10, 25, 20, 10, 10, 12, 15, 12, 20, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'配件申请报表_{timestamp}.xlsx'
    
    elif report_type == 'transfers':
        # 设备调拨报表
        ws = wb.create_sheet("设备调拨")
        headers = ['调拨ID', '设备名称', '调出部门', '调入部门', '申请人', 
                  '状态', '创建时间', '审批时间', '原因']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        transfers = EquipmentTransfer.query.order_by(EquipmentTransfer.created_date.desc()).all()
        for t in transfers:
            ws.append([
                t.id,
                t.equipment.name if t.equipment else '',
                t.from_department or '',
                t.to_department or '',
                t.requester.real_name or t.requester.username if t.requester else '',
                t.status,
                t.created_date.strftime('%Y-%m-%d %H:%M:%S') if t.created_date else '',
                t.approved_date.strftime('%Y-%m-%d %H:%M:%S') if hasattr(t, 'approved_date') and t.approved_date else '',
                t.reason or ''
            ])
        
        column_widths = [10, 25, 15, 15, 12, 12, 20, 20, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'设备调拨报表_{timestamp}.xlsx'
    
    elif report_type == 'asset_lifecycle':
        # 资产生命周期报表
        ws = wb.create_sheet("资产生命周期")
        headers = ['事件ID', '设备名称', '事件类型', '事件日期', '涉及成本', 
                  '负责人', '备注']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        events = AssetLifecycle.query.order_by(AssetLifecycle.event_date.desc()).all()
        for event in events:
            event_type_map = {
                'purchase': '采购',
                'deployment': '部署',
                'maintenance': '维护',
                'repair': '维修',
                'upgrade': '升级',
                'transfer': '调拨',
                'scrap': '报废'
            }
            
            ws.append([
                event.id,
                event.equipment.name if event.equipment else '',
                event_type_map.get(event.event_type, event.event_type),
                event.event_date.strftime('%Y-%m-%d') if event.event_date else '',
                float(event.cost_involved or 0),
                event.responsible_user.real_name or event.responsible_user.username if event.responsible_user else '',
                event.notes or '' if hasattr(event, 'notes') else (event.description or '')
            ])
        
        column_widths = [10, 25, 12, 12, 12, 12, 35]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'资产生命周期报表_{timestamp}.xlsx'
    
    elif report_type == 'scraps':
        # 设备报废报表
        ws = wb.create_sheet("设备报废")
        headers = ['报废ID', '设备名称', '设备编号', '设备类型', '所属部门', 
                  '申请人', '报废原因', '状态', '申请时间', '审批时间']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        scraps = EquipmentScrap.query.order_by(EquipmentScrap.created_date.desc()).all()
        for s in scraps:
            status_map = {
                'submitted': '已提交',
                'pending': '待审批',
                'approved': '已批准',
                'rejected': '已拒绝',
                'cancelled': '已取消'
            }
            ws.append([
                s.id,
                s.equipment.name if s.equipment else '',
                s.equipment.equipment_number if s.equipment else '',
                s.equipment.equipment_type.name if s.equipment and s.equipment.equipment_type else '',
                (s.equipment.department if isinstance(s.equipment.department, str) else getattr(s.equipment.department, 'name', '')) if s.equipment and s.equipment.department else '',
                s.requester.username if s.requester else '',
                s.description or '',
                status_map.get(s.status, s.status),
                s.created_date.strftime('%Y-%m-%d %H:%M:%S') if s.created_date else '',
                s.updated_date.strftime('%Y-%m-%d %H:%M:%S') if s.status == 'approved' and s.updated_date else ''
            ])
        
        column_widths = [10, 25, 20, 15, 15, 12, 35, 12, 20, 20]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'设备报废报表_{timestamp}.xlsx'
    
    elif report_type == 'loans':
        # 设备借用报表
        ws = wb.create_sheet("设备借用")
        headers = ['借用ID', '设备名称', '设备编号', '借用人', '借用部门', 
                  '计划开始', '计划结束', '实际借出', '实际归还', '状态', '备注']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        loans = EquipmentLoan.query.order_by(EquipmentLoan.created_date.desc()).all()
        for loan in loans:
            status_map = {
                'submitted': '已提交',
                'approved': '已批准',
                'borrowed': '已借出',
                'returned': '已归还',
                'rejected': '已拒绝',
                'cancelled': '已取消'
            }
            ws.append([
                loan.id,
                loan.equipment.name if loan.equipment else '',
                loan.equipment.equipment_number if loan.equipment else '',
                loan.requester.username if loan.requester else '',
                loan.requester_dept or '',
                loan.start_date.strftime('%Y-%m-%d') if loan.start_date else '',
                loan.end_date.strftime('%Y-%m-%d') if loan.end_date else '',
                loan.borrowed_date.strftime('%Y-%m-%d %H:%M') if loan.borrowed_date else '',
                loan.returned_date.strftime('%Y-%m-%d %H:%M') if loan.returned_date else '',
                status_map.get(loan.status, loan.status),
                loan.notes or ''
            ])
        
        column_widths = [10, 25, 20, 12, 15, 12, 12, 18, 18, 12, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        filename = f'设备借用报表_{timestamp}.xlsx'
    
    else:
        # 综合汇总报表
        ws = wb.create_sheet("综合统计")
        headers = ['统计项', '数量/金额']
        ws.append(headers)
        
        for col in range(1, 3):
            cell = ws.cell(1, col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # 基础统计
        stats_data = [
            ['用户总数', User.query.count()],
            ['部门总数', Department.query.count()],
            ['设备总数', Equipment.query.count()],
            ['设备类型数', EquipmentType.query.count()],
            ['维修工单总数', RepairOrder.query.count()],
            ['配件申请总数', PartRequestOrder.query.count()],
            ['设备调拨总数', EquipmentTransfer.query.count()],
            ['设备借用总数', EquipmentLoan.query.count()],
            ['设备报废总数', EquipmentScrap.query.count()],
            ['', ''],
            ['设备总价值', float(db.session.query(func.sum(Equipment.price)).scalar() or 0)],
            ['维修总成本', float(db.session.query(func.sum(AssetCost.maintenance_cost)).scalar() or 0)],
        ]
        
        for row_data in stats_data:
            ws.append(row_data)
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
        
        filename = f'综合统计报表_{timestamp}.xlsx'
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
